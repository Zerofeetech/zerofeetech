import os, sqlite3, secrets, io, json, urllib.request, urllib.error, base64
from pathlib import Path
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, abort
from werkzeug.security import generate_password_hash, check_password_hash

BASE=Path(__file__).resolve().parent
GOOGLE_API_KEY=os.environ.get('GOOGLE_API_KEY','').strip()
AI_MODEL=os.environ.get('AI_MODEL','gemini-2.5-flash').strip()
DB_PATH=os.environ.get('DATABASE_PATH',str(BASE/'zerofeetech.db'))
ADMIN_EMAIL=os.environ.get('ADMIN_EMAIL','admin@zerofeetech.com').strip().lower()
CONTACT_EMAIL=os.environ.get('CONTACT_EMAIL','zerofeetech@gmail.com').strip()
WHATSAPP_NUMBER=''.join(ch for ch in os.environ.get('WHATSAPP_NUMBER','') if ch.isdigit())
ADMIN_PASSWORD=os.environ.get('ADMIN_PASSWORD','')
app=Flask(__name__)
app.secret_key=os.environ.get('SECRET_KEY') or secrets.token_hex(32)
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax', SESSION_COOKIE_SECURE=os.environ.get('COOKIE_SECURE','true').lower()=='true', MAX_CONTENT_LENGTH=12*1024*1024)

def connect():
    c=sqlite3.connect(DB_PATH,timeout=15); c.row_factory=sqlite3.Row; return c

def init_db():
    Path(DB_PATH).parent.mkdir(parents=True,exist_ok=True)
    with connect() as c:
        c.executescript('''CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,email TEXT UNIQUE NOT NULL,password_hash TEXT NOT NULL,role TEXT NOT NULL DEFAULT 'user',created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS projects(id INTEGER PRIMARY KEY AUTOINCREMENT,user_id INTEGER,title TEXT NOT NULL,category TEXT NOT NULL,details TEXT NOT NULL,status TEXT NOT NULL DEFAULT 'New',created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS saved_items(id INTEGER PRIMARY KEY AUTOINCREMENT,user_id INTEGER NOT NULL,title TEXT NOT NULL,content TEXT NOT NULL,created_at TEXT DEFAULT CURRENT_TIMESTAMP);''')

def csrf():
    if '_csrf' not in session: session['_csrf']=secrets.token_urlsafe(24)
    return session['_csrf']
app.jinja_env.globals['csrf_token']=csrf

def me():
    uid=session.get('uid')
    if not uid:return None
    with connect() as c:r=c.execute('SELECT id,name,email,role,created_at FROM users WHERE id=?',(uid,)).fetchone()
    return dict(r) if r else None
app.jinja_env.globals['current_user']=me
app.jinja_env.globals['contact_email']=CONTACT_EMAIL
app.jinja_env.globals['whatsapp_number']=WHATSAPP_NUMBER

@app.before_request
def check_csrf():
    if request.method=='POST':
        sent=request.form.get('_csrf') or request.headers.get('X-CSRF-Token')
        if not sent or not secrets.compare_digest(sent,session.get('_csrf','')): abort(400,'Security token expired. Refresh and try again.')

def login_required(f):
    @wraps(f)
    def w(*a,**kw):
        if not me(): flash('Sign in to access your workspace.','info'); return redirect(url_for('login'))
        return f(*a,**kw)
    return w

def admin_required(f):
    @wraps(f)
    def w(*a,**kw):
        u=me()
        if not u or u['role']!='admin': abort(403)
        return f(*a,**kw)
    return w

@app.get('/')
def home():return render_template('index.html',title='ZeroFeeTech')
@app.get('/tools')
def tools():return render_template('tools.html',title='All tools')

ALLOWED_IMAGE_EXTS={'.png','.jpg','.jpeg','.webp'}

def extract_upload(file):
    # Extract readable text from supported documents without saving uploads to disk.
    from pathlib import PurePath
    name=(file.filename or 'uploaded-file').replace('\\','/').split('/')[-1][:180]
    ext=PurePath(name).suffix.lower()
    raw=file.read()
    if not raw:
        return name, '', None
    if ext in {'.txt','.md','.csv','.json'}:
        return name, raw.decode('utf-8-sig',errors='replace')[:60000], None
    if ext=='.pdf':
        try:
            from pypdf import PdfReader
            reader=PdfReader(io.BytesIO(raw))
            pages=[]
            for i,page in enumerate(reader.pages[:80]):
                pages.append(f'\n[Page {i+1}]\n'+(page.extract_text() or ''))
            return name, '\n'.join(pages)[:60000], None
        except Exception:
            return name, '', 'Could not read this PDF. It may be scanned or damaged; try a text-based PDF or paste the text.'
    if ext=='.docx':
        try:
            from docx import Document
            doc=Document(io.BytesIO(raw))
            text='\n'.join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                text+='\n'+'\n'.join(' | '.join(cell.text for cell in row.cells) for row in table.rows)
            return name, text[:60000], None
        except Exception:
            return name, '', 'Could not read this DOCX. Please re-save it as a standard .docx file.'
    if ext=='.pptx':
        try:
            from pptx import Presentation
            deck=Presentation(io.BytesIO(raw)); slides=[]
            for i,slide in enumerate(deck.slides[:80]):
                bits=[]
                for shape in slide.shapes:
                    if getattr(shape,'has_text_frame',False) and shape.text.strip(): bits.append(shape.text.strip())
                    if getattr(shape,'has_table',False):
                        bits.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
                slides.append(f'\n[Slide {i+1}]\n'+'\n'.join(bits))
            return name, '\n'.join(slides)[:60000], None
        except Exception:
            return name, '', 'Could not read this PPTX. Please re-save it as a standard .pptx file.'
    if ext in ALLOWED_IMAGE_EXTS:
        mime={'.jpg':'image/jpeg','.jpeg':'image/jpeg','.png':'image/png','.webp':'image/webp'}.get(ext)
        if mime:
            return name, '', {'mime_type':mime,'data':base64.b64encode(raw).decode('ascii')}
    return name, '', 'Unsupported file type. Use PDF, DOCX, PPTX, TXT, MD, CSV, JSON, PNG, JPG or WEBP.'

def call_gemini(prompt, image_parts=None):
    if not GOOGLE_API_KEY:
        return None, 'AI is not connected yet. Add GOOGLE_API_KEY in Render Environment to enable model-powered answers. This page will not pretend a template is a genuine AI answer.'
    body_parts=[{'text':prompt}]
    for part in (image_parts or []):
        body_parts.append({'inline_data':{'mime_type':part['mime_type'],'data':part['data']}})
    payload={'contents':[{'role':'user','parts':body_parts}], 'generationConfig':{'temperature':0.25,'maxOutputTokens':7000}}
    endpoint=f'https://generativelanguage.googleapis.com/v1beta/models/{AI_MODEL}:generateContent?key={GOOGLE_API_KEY}'
    req=urllib.request.Request(endpoint,data=json.dumps(payload).encode(),headers={'Content-Type':'application/json'},method='POST')
    try:
        with urllib.request.urlopen(req,timeout=50) as response: data=json.loads(response.read().decode('utf-8'))
        answer='\n'.join(part.get('text','') for candidate in data.get('candidates',[]) for part in candidate.get('content',{}).get('parts',[]) if part.get('text'))
        if not answer: return None, 'The AI service returned no text. Try a smaller file or a more specific question.'
        return answer, None
    except urllib.error.HTTPError as e:
        detail=''
        try: detail=e.read().decode('utf-8')[:500]
        except Exception: pass
        return None, f'AI provider error ({e.code}). Check the API key, model access and free-tier quota. {detail}'
    except Exception:
        return None, 'Could not reach the AI provider. Please try again in a moment.'

@app.get('/assignment')
def assignment():
    return render_template('assignment.html',title='Assignment Studio',ai_connected=bool(GOOGLE_API_KEY))

@app.post('/assignment')
def assignment_post():
    task=request.form.get('task','explain').strip().lower()
    if task not in {'explain','solve','outline','review','summarize','quiz','citation'}: task='explain'
    question=request.form.get('question','').strip()[:12000]
    level=request.form.get('level','undergraduate').strip()[:80]
    language=request.form.get('language','English').strip()[:60]
    files=[f for f in request.files.getlist('files') if f and f.filename]
    if not question and not files:
        flash('Type your question or upload at least one file first.','error')
        return redirect(url_for('assignment'))
    if len(files)>8:
        flash('Please upload up to 8 files at a time.','error')
        return redirect(url_for('assignment'))
    docs=[]; image_parts=[]; warnings=[]
    for file in files:
        name,content,image_or_error=extract_upload(file)
        if isinstance(image_or_error,dict): image_parts.append(image_or_error); docs.append((name,'[Image attached for visual analysis]'))
        elif isinstance(image_or_error,str) and image_or_error: warnings.append(f'{name}: {image_or_error}')
        else: docs.append((name,content or '[No selectable text found in this file]'))
    context='\n\n'.join(f'===== FILE: {name} =====\n{text}' for name,text in docs)[:100000]
    task_guides={
      'explain':'Explain the material clearly from fundamentals to advanced level, define key terms, show examples, and identify ambiguity.',
      'solve':'Work through the supplied questions step by step. Show reasoning, formulas or evidence where relevant, then provide a concise final answer and self-check. If information is missing, state assumptions instead of inventing facts.',
      'outline':'Create a detailed assignment outline with thesis or objective, section headings, key points, evidence needed, and a completion checklist. Do not fabricate sources.',
      'review':"Review the student's own draft for accuracy, structure, clarity, grammar, unsupported claims, and missing citations. Give specific suggestions and a revised sample passage where helpful.",
      'summarize':'Summarize the material in plain language, then give key points, definitions, and a short revision sheet.',
      'quiz':'Create a study quiz from the provided material with mixed difficulty and an answer key with explanations at the end.',
      'citation':'Identify claims that need sources and explain what source types to look for. Format only bibliographic details that are actually provided; never invent citations or page numbers.'}
    prompt=(f"You are ZeroFeeTech Study & Assignment Assistant. Support learning, accuracy and student independence. Respond in {language}. Student level: {level}. Task: {task}.\n\n"
      f"INSTRUCTIONS: {task_guides[task]}\n\n"
      "Important: Never claim to bypass AI detectors or disguise authorship. Help the student understand the subject and create a draft they must verify, personalize and submit according to their institution's rules. Do not fabricate references, quotes, facts, page numbers, data, or answers. If a file is unreadable or a question is ambiguous, say so. For calculations, show steps and verify the result. Structure the response with clear headings and actionable next steps.\n\n"
      f"STUDENT QUESTION / INSTRUCTIONS:\n{question or '(No extra question; infer the task from uploaded files and state any uncertainty.)'}\n\n"
      f"EXTRACTED FILE CONTENT:\n{context or '(No text extracted from uploaded files.)'}")
    answer,error=call_gemini(prompt,image_parts)
    result={'answer':answer,'error':error,'warnings':warnings,'files':[n for n,_ in docs], 'question':question,'task':task}
    if answer and me():
        try:
            with connect() as c:c.execute('INSERT INTO saved_items(user_id,title,content) VALUES(?,?,?)',(me()['id'],f'Assignment Studio — {task.title()}',answer[:20000]))
            result['saved']=True
        except Exception: result['saved']=False
    return render_template('assignment.html',title='Assignment Studio',ai_connected=bool(GOOGLE_API_KEY),result=result)

@app.get('/signup')
def signup():return render_template('auth.html',mode='signup',title='Create account')
@app.post('/signup')
def signup_post():
    name=request.form.get('name','').strip()[:100]; email=request.form.get('email','').strip().lower()[:254]; pw=request.form.get('password','')
    if len(name)<2 or '@' not in email or len(pw)<10: flash('Use a valid email and a password of at least 10 characters.','error'); return redirect(url_for('signup'))
    role='admin' if email==ADMIN_EMAIL and ADMIN_PASSWORD and secrets.compare_digest(pw,ADMIN_PASSWORD) else 'user'
    try:
        with connect() as c:c.execute('INSERT INTO users(name,email,password_hash,role) VALUES(?,?,?,?)',(name,email,generate_password_hash(pw),role))
        flash('Account created. Sign in to continue.','success');return redirect(url_for('login'))
    except sqlite3.IntegrityError:flash('An account with that email already exists. Try signing in.','error');return redirect(url_for('signup'))
@app.get('/login')
def login():return render_template('auth.html',mode='login',title='Sign in')
@app.post('/login')
def login_post():
    email=request.form.get('email','').strip().lower(); pw=request.form.get('password','')
    with connect() as c:r=c.execute('SELECT * FROM users WHERE email=?',(email,)).fetchone()
    # Safe recovery path for the configured administrator: the hosting environment variable is authoritative.
    if email==ADMIN_EMAIL and ADMIN_PASSWORD and secrets.compare_digest(pw,ADMIN_PASSWORD):
        with connect() as c:
            if r:
                c.execute("UPDATE users SET role='admin',password_hash=? WHERE id=?",(generate_password_hash(pw),r['id']))
                r=c.execute('SELECT * FROM users WHERE id=?',(r['id'],)).fetchone()
            else:
                c.execute("INSERT INTO users(name,email,password_hash,role) VALUES(?,?,?,'admin')",('ZeroFeeTech Admin',ADMIN_EMAIL,generate_password_hash(pw)))
                r=c.execute('SELECT * FROM users WHERE email=?',(ADMIN_EMAIL,)).fetchone()
    if not r or not check_password_hash(r['password_hash'],pw):flash('Email or password is incorrect.','error');return redirect(url_for('login'))
    session.clear();session['uid']=r['id'];session['_csrf']=secrets.token_urlsafe(24)
    return redirect(url_for('admin' if r['role']=='admin' else 'dashboard'))
@app.post('/logout')
def logout():session.clear();flash('Signed out successfully.','success');return redirect(url_for('home'))
@app.get('/dashboard')
@login_required
def dashboard():
    u=me()
    with connect() as c:
        projects=c.execute('SELECT * FROM projects WHERE user_id=? ORDER BY id DESC',(u['id'],)).fetchall();saved=c.execute('SELECT * FROM saved_items WHERE user_id=? ORDER BY id DESC',(u['id'],)).fetchall()
    return render_template('dashboard.html',title='Workspace',projects=projects,saved=saved)
@app.post('/save')
@login_required
def save():
    title=request.form.get('title','Saved prompt').strip()[:120];content=request.form.get('content','').strip()[:20000]
    if content:
        with connect() as c:c.execute('INSERT INTO saved_items(user_id,title,content) VALUES(?,?,?)',(me()['id'],title,content))
        flash('Saved to your library.','success')
    else:flash('There is no content to save yet.','error')
    return redirect(url_for('dashboard'))
@app.post('/request-project')
def request_project():
    title=request.form.get('title','').strip()[:160];category=request.form.get('category','Website').strip()[:80];details=request.form.get('details','').strip()[:10000];u=me()
    if len(title)<3 or len(details)<10:flash('Please add a title and a little more project detail.','error');return redirect(url_for('home')+'#start')
    with connect() as c:c.execute('INSERT INTO projects(user_id,title,category,details) VALUES(?,?,?,?)',(u['id'] if u else None,title,category,details))
    flash('Project brief received.','success');return redirect(url_for('dashboard') if u else url_for('home')+'#start')
@app.post('/api/build-prompt')
def build_prompt():
    d=request.get_json(silent=True) or {};category=str(d.get('category','Website'))[:80];idea=str(d.get('idea','')).strip()[:3000];audience=str(d.get('audience','general users'))[:300];style=str(d.get('style','premium, accessible, modern'))[:300]
    if len(idea)<5:return jsonify(error='Please describe your idea in a little more detail.'),400
    prompt=f'''Act as a senior product designer, software architect, and quality engineer.\n\nPROJECT TYPE: {category}\nCORE IDEA: {idea}\nTARGET USERS: {audience}\nDESIGN DIRECTION: {style}\n\nDELIVERABLES\n1. Product objective, assumptions, and user stories.\n2. Information architecture and end-to-end user journeys.\n3. A distinctive responsive design system and accessibility requirements.\n4. Maintainable technical architecture and implementation steps.\n5. Loading, empty, validation, error, and success states.\n6. Testing plan for mobile, desktop, security, performance, and accessibility.\n7. Deployment, environment-variable, monitoring, backup, and rollback checklist.\n\nBe specific. Do not claim integrations are live unless they are implemented and tested. Identify external services and API keys required.'''
    return jsonify(prompt=prompt)
@app.get('/admin')
@admin_required
def admin():
    with connect() as c:
        stats={'users':c.execute('SELECT COUNT(*) FROM users').fetchone()[0],'projects':c.execute('SELECT COUNT(*) FROM projects').fetchone()[0],'new':c.execute("SELECT COUNT(*) FROM projects WHERE status='New'").fetchone()[0],'saved':c.execute('SELECT COUNT(*) FROM saved_items').fetchone()[0]}
        users=c.execute('SELECT id,name,email,role,created_at FROM users ORDER BY id DESC LIMIT 200').fetchall();projects=c.execute('SELECT p.*,u.name AS owner FROM projects p LEFT JOIN users u ON u.id=p.user_id ORDER BY p.id DESC LIMIT 200').fetchall()
    return render_template('admin.html',title='Admin control center',stats=stats,users=users,projects=projects)
@app.post('/admin/project/<int:pid>/status')
@admin_required
def status(pid):
    s=request.form.get('status','New')
    if s not in {'New','Reviewing','In progress','Completed','Declined'}:abort(400)
    with connect() as c:c.execute('UPDATE projects SET status=? WHERE id=?',(s,pid))
    flash('Project status updated.','success');return redirect(url_for('admin')+'#projects')
@app.post('/admin/user/<int:uid>/role')
@admin_required
def role(uid):
    r=request.form.get('role','user')
    if r not in {'user','admin'}:abort(400)
    if uid==me()['id'] and r!='admin':flash('You cannot remove your own admin role here.','error')
    else:
        with connect() as c:c.execute('UPDATE users SET role=? WHERE id=?',(r,uid))
        flash('User role updated.','success')
    return redirect(url_for('admin')+'#users')
@app.errorhandler(400)
def bad(e):return render_template('error.html',title='Request error',code=400,message=getattr(e,'description','Refresh and try again.')),400
@app.errorhandler(403)
def forbidden(e):return render_template('error.html',title='Access restricted',code=403,message='This area is restricted to authorized administrators.'),403
@app.errorhandler(404)
def missing(e):return render_template('error.html',title='Page not found',code=404,message='The page could not be found.'),404
init_db()
if __name__=='__main__':app.run(host='0.0.0.0',port=int(os.environ.get('PORT',5000)),debug=False)
